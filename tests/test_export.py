from pathlib import Path
import sqlite3

import pytest
from openpyxl import load_workbook

from csindex_local.db import Database
from csindex_local.models import IndexRecord, VolatilitySnapshot, YieldSnapshot


@pytest.fixture
def seed_database(tmp_path: Path) -> Database:
    database = Database(tmp_path / "csindex.db")
    database.initialize()
    database.upsert_indices(
        [
            IndexRecord("000300", "沪深300", "是", {}),
            IndexRecord("000905", "中证500", "是", {}),
            IndexRecord("000852", "中证1000", "是", {}),
        ]
    )
    database.create_or_get_scope("scope-1", ["000300", "000905", "000852"])

    # An older, larger value proves latest_rows does not return historical data.
    database.upsert_yield(
        YieldSnapshot("000300", "2026-09-02", 10, 10, 10, 99, 10, 10)
    )
    database.merge_volatility(
        VolatilitySnapshot("000300", "2026-09-02", 10, 10, 10)
    )
    database.upsert_yield(
        YieldSnapshot("000300", "2026-09-03", 1.234, -2.345, 3, 4, 5, 6)
    )
    database.merge_volatility(
        VolatilitySnapshot("000300", "2026-09-03", 7, 8, 9)
    )
    database.upsert_yield(
        YieldSnapshot("000905", "2026-09-03", -1, 2, 3, 8, 5, 6)
    )
    database.merge_volatility(
        VolatilitySnapshot("000905", "2026-09-03", 7, 8, 9)
    )
    database.upsert_yield(
        YieldSnapshot("000852", "2026-09-03", 1, None, 3, None, 5, 6)
    )
    database.merge_volatility(
        VolatilitySnapshot("000852", "2026-09-03", 7, 8, 9)
    )
    return database


@pytest.fixture
def query_service(seed_database: Database):
    from csindex_local.query_service import QueryService

    return QueryService(seed_database)


def test_latest_rows_return_one_newest_snapshot_per_scope_member(query_service):
    rows = query_service.latest_rows("scope-1")

    assert [row.index_code for row in rows] == ["000300", "000905", "000852"]
    assert rows[0].data_date == "2026-09-03"
    assert rows[0].one_year == 4.0
    assert rows[2].missing_count == 2
    assert rows[2].is_complete is False


def test_latest_rows_retains_scope_members_without_a_snapshot(seed_database: Database):
    from csindex_local.query_service import QueryService

    seed_database.upsert_indices([IndexRecord("000001", "上证指数", "是", {})])
    seed_database.create_or_get_scope("empty-snapshot", ["000001"])

    row = QueryService(seed_database).latest_rows("empty-snapshot")[0]
    assert row.index_code == "000001"
    assert row.data_date == "--"
    assert row.one_year is None
    assert row.missing_count == 9
    assert row.is_complete is False


def test_ranking_excludes_missing_one_year_and_sorts_numeric_descending(query_service):
    rows = query_service.rank(query_service.latest_rows("scope-1"), "one_year", None)

    assert all(row.one_year is not None for row in rows)
    assert [row.one_year for row in rows] == [8.0, 4.0]
    assert [row.index_code for row in query_service.rank(rows, "one_year", 1)] == ["000905"]
    with pytest.raises(ValueError, match="unsupported ranking field"):
        query_service.rank(rows, "not_a_metric", None)


def test_export_has_four_sheets_and_typed_styled_cells(
    seed_database: Database, tmp_path: Path
):
    from csindex_local.excel_exporter import ExcelExporter

    path = ExcelExporter(seed_database).export("scope-1", tmp_path / "out.xlsx", 100)
    workbook = load_workbook(path, data_only=False)

    assert workbook.sheetnames == [
        "完整九项指标",
        "近一年收益率排名",
        "缺失与异常",
        "说明",
    ]
    sheet = workbook["完整九项指标"]
    assert sheet["B3"].value == "000300"
    assert sheet["B3"].data_type == "s"
    assert sheet["B3"].number_format == "@"
    assert sheet["D3"].value == pytest.approx(1.234)
    assert sheet["D3"].number_format == "0.00"
    assert sheet["E3"].font.color.rgb.endswith("008000")
    assert sheet["D3"].font.color.rgb.endswith("C00000")
    assert sheet["E5"].value == "--"
    assert sheet.freeze_panes == "A3"
    assert sheet.auto_filter.ref == "A2:O5"
    assert workbook["近一年收益率排名"]["D3"].value == 8.0
    assert workbook["缺失与异常"]["A3"].number_format == "@"
    assert "https://www.csindex.com.cn" in [
        row[1] for row in workbook["说明"].iter_rows(values_only=True) if len(row) > 1
    ]


def test_export_uses_timestamped_path_when_target_is_locked(
    seed_database: Database, tmp_path: Path, monkeypatch
):
    from openpyxl import Workbook

    from csindex_local.excel_exporter import ExcelExporter

    target = tmp_path / "occupied.xlsx"
    target.write_bytes(b"original workbook remains untouched")
    original_save = Workbook.save

    def save_with_locked_target(workbook, filename):
        if Path(filename) == target:
            raise PermissionError("occupied by Excel")
        return original_save(workbook, filename)

    monkeypatch.setattr(Workbook, "save", save_with_locked_target)
    exported = ExcelExporter(seed_database).export("scope-1", target, None)

    assert exported != target
    assert exported.exists()
    assert target.read_bytes() == b"original workbook remains untouched"


def test_export_marks_volatility_date_mismatch_as_an_exception(
    seed_database: Database, tmp_path: Path
):
    from csindex_local.excel_exporter import ExcelExporter

    seed_database.create_or_get_scope("fixed:1", ["000905"])
    with sqlite3.connect(seed_database.path) as connection:
        connection.execute(
            """
            INSERT INTO crawl_runs (
                id, scope_type, scope_value, target_data_date, status, started_at,
                total_tasks
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                "run-date-mismatch",
                "fixed",
                "1",
                "2026-09-03",
                "completed_with_failures",
                "2026-09-03T00:00:00+00:00",
                1,
            ),
        )
        connection.execute(
            """
            INSERT INTO crawl_tasks (
                run_id, index_code, endpoint, target_data_date, status, attempts,
                available_at, last_error, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                "run-date-mismatch",
                "000905",
                "volatility",
                "2026-09-03",
                "failed",
                1,
                "2026-09-03T00:00:00+00:00",
                "volatility response returned a different data date",
                "2026-09-03T00:00:00+00:00",
            ),
        )

    path = ExcelExporter(seed_database).export("fixed:1", tmp_path / "out.xlsx", 100)
    issues = load_workbook(path, data_only=False)["缺失与异常"]

    assert issues["D3"].value == "数据日期不一致"
    assert issues["D3"].value != "请求失败"
