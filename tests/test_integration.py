"""Offline, loopback end-to-end coverage for the complete crawl stack."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from csindex_local.config import AppConfig
from csindex_local.crawler import CrawlControl, Crawler
from csindex_local.csindex_client import CsindexClient
from csindex_local.db import Database
from csindex_local.excel_exporter import ExcelExporter
from csindex_local.models import ScopeSelection, UpdateMode

from fake_csindex_server import FakeCsindexServer


TARGET_DATE = "2026-09-03"


class VirtualClock:
    """A clock whose sleeps move time forward instead of waiting."""

    def __init__(self) -> None:
        self.current = datetime(2026, 9, 4, tzinfo=timezone.utc)
        self.sleeps: list[float] = []

    def now(self) -> datetime:
        return self.current

    def sleep(self, seconds: float) -> None:
        self.sleeps.append(seconds)
        self.current += timedelta(seconds=seconds)

    def advance(self, seconds: float) -> None:
        self.current += timedelta(seconds=seconds)


@pytest.fixture
def fake_server() -> FakeCsindexServer:
    with FakeCsindexServer(index_count=20) as server:
        yield server


def no_wait_config() -> AppConfig:
    return AppConfig(
        data_dir="",
        export_dir="",
        request_delay_min_seconds=0,
        request_delay_max_seconds=0,
        batch_size=100,
        batch_rest_seconds=0,
    )


def make_crawler(
    server: FakeCsindexServer, database: Database, clock: VirtualClock, *, timeout: float = 1
) -> Crawler:
    return Crawler(
        CsindexClient(server.base_url, timeout_seconds=timeout),
        database,
        config=no_wait_config(),
        clock=clock.now,
        sleep=clock.sleep,
        random_uniform=lambda minimum, maximum: 0,
    )


def new_database(tmp_path: Path) -> Database:
    database = Database(tmp_path / "data" / "csindex.db")
    database.initialize()
    return database


def task_rows(database: Database, run_id: str):
    with database._connection() as connection:
        return connection.execute(
            "SELECT * FROM crawl_tasks WHERE run_id = ? ORDER BY id", (run_id,)
        ).fetchall()


def count_complete(database: Database) -> int:
    with database._connection() as connection:
        return int(
            connection.execute(
                "SELECT COUNT(*) AS count FROM metric_snapshots WHERE is_complete = 1"
            ).fetchone()["count"]
        )


def test_twenty_indices_resume_same_database_without_repeating_successes(
    fake_server: FakeCsindexServer, tmp_path: Path
) -> None:
    """An explicit interruption after 13 complete indices resumes all 40 tasks."""

    database = new_database(tmp_path)
    clock = VirtualClock()
    codes = tuple(f"{number:06d}" for number in range(1, 21))
    first = make_crawler(fake_server, database, clock)
    run_id = first.prepare_run(ScopeSelection("codes", codes), UpdateMode.FORCE)

    # Queue ordering is all yields then all volatilities, so 33 tasks are needed
    # to reach exactly 13 complete nine-metric snapshots before interruption.
    for _ in range(33):
        first.run_one(run_id)
    assert count_complete(database) == 13

    restarted = make_crawler(fake_server, database, clock)
    progress = restarted.run(run_id, CrawlControl(), lambda event: None)
    output = ExcelExporter(database).export(
        database_scope_id(database),
        tmp_path / "exports" / "result.xlsx",
        None,
    )

    assert progress.success_tasks == 40
    assert progress.failed_tasks == 0
    assert count_complete(database) == 20
    assert output.exists()
    assert load_workbook(output).sheetnames == [
        "完整九项指标",
        "近一年收益率排名",
        "缺失与异常",
        "说明",
    ]
    with database._connection() as connection:
        persisted = connection.execute(
            "SELECT COUNT(*) AS count, SUM(is_success) AS successes FROM raw_responses"
        ).fetchone()
    # One catalogue page, one CSI-300 probe, and forty detail requests all
    # travelled through the real client observer into SQLite.
    assert (persisted["count"], persisted["successes"]) == (42, 42)
    for code in codes:
        assert fake_server.detail_request_counts[f"yield:{code}"] == 1
        assert fake_server.detail_request_counts[f"volatility:{code}"] == 1


def database_scope_id(database: Database) -> str:
    """Read the persisted codes-scope ID; test code never duplicates app logic."""

    with database._connection() as connection:
        row = connection.execute(
            "SELECT id FROM crawl_scopes WHERE scope_type = 'codes' ORDER BY created_at LIMIT 1"
        ).fetchone()
    assert row is not None
    return str(row["id"])


def test_waf_cooldown_is_persisted_before_restart_and_no_task_follows_block(
    fake_server: FakeCsindexServer, tmp_path: Path
) -> None:
    database = new_database(tmp_path)
    clock = VirtualClock()
    fake_server.set_indices(("000001", "000002"))
    fake_server.plan("yield", "000001", "waf_403")
    crawler = make_crawler(fake_server, database, clock)
    run_id = crawler.prepare_run(
        ScopeSelection("codes", ("000001", "000002")), UpdateMode.FORCE
    )

    waiting = crawler.run(run_id, CrawlControl(), lambda event: None)

    assert waiting.pending_tasks == 4
    assert dict(fake_server.detail_request_counts) == {
        "yield:000300": 1,
        "yield:000001": 1,
    }
    state = database.get_runtime_state("waf_cooldown")
    assert state is not None
    assert state["next_cooldown_seconds"] == 3600
    with database._connection() as connection:
        blocked = connection.execute(
            "SELECT http_status, payload, is_success FROM raw_responses WHERE http_status = 403"
        ).fetchone()
    assert blocked is not None
    assert "WAF" in blocked["payload"]
    assert blocked["is_success"] == 0

    # A fresh coordinator reads the same persisted deadline.  Its injected
    # sleep advances virtual time, so no real 30-minute wait occurs.
    restarted = make_crawler(fake_server, database, clock)
    finished = restarted.run(run_id, CrawlControl(), lambda event: None)

    assert 1800 in clock.sleeps
    assert finished.success_tasks == 4
    assert database.get_runtime_state("waf_cooldown") is None
    assert fake_server.detail_request_counts["yield:000002"] == 1


def test_business_404_fails_only_its_request_and_other_indices_continue(
    fake_server: FakeCsindexServer, tmp_path: Path
) -> None:
    database = new_database(tmp_path)
    clock = VirtualClock()
    fake_server.set_indices(("000001", "000002"))
    fake_server.plan("yield", "000001", "business_404")
    crawler = make_crawler(fake_server, database, clock)
    run_id = crawler.prepare_run(
        ScopeSelection("codes", ("000001", "000002")), UpdateMode.FORCE
    )

    progress = crawler.run(run_id, CrawlControl(), lambda event: None)

    failed = [row for row in task_rows(database, run_id) if row["status"] == "failed"]
    assert progress.success_tasks == 2
    assert progress.failed_tasks == 2  # the dependent volatility task cannot run
    assert fake_server.detail_request_counts["yield:000001"] == 1
    assert fake_server.detail_request_counts["yield:000002"] == 1
    assert fake_server.detail_request_counts["volatility:000002"] == 1
    assert "HTTP 404" in failed[0]["last_error"]
    assert database.get_runtime_state("waf_cooldown") is None
    with database._connection() as connection:
        missing = connection.execute(
            "SELECT payload, is_success FROM raw_responses WHERE http_status = 404"
        ).fetchone()
    assert missing is not None
    assert "does not exist" in missing["payload"]
    assert missing["is_success"] == 0


def test_timeout_bad_json_and_date_mismatch_retry_or_export_anomalies(
    fake_server: FakeCsindexServer, tmp_path: Path
) -> None:
    database = new_database(tmp_path)
    clock = VirtualClock()
    fake_server.set_indices(("000001", "000002", "000003"))
    fake_server.plan("yield", "000001", "timeout", "ok")
    fake_server.plan("yield", "000002", "bad_json", "ok")
    fake_server.plan("yield", "000003", ("yield_date", "2026-09-02"))
    # Keep catalogue/probe setup tolerant of a busy Windows test runner; only
    # the planned detail request needs the deliberately tiny transport timeout.
    crawler = make_crawler(fake_server, database, clock, timeout=1)
    run_id = crawler.prepare_run(
        ScopeSelection("codes", ("000001", "000002", "000003")), UpdateMode.FORCE
    )
    crawler._client.timeout_seconds = 0.03

    # Drive the three yields without allowing a volatility dependency wait to
    # end the run loop before every injected transport scenario is exercised.
    for _ in range(4):
        first = crawler.run_one(run_id)
    assert first.pending_tasks == 5
    assert fake_server.detail_request_counts["yield:000001"] == 1
    assert fake_server.detail_request_counts["yield:000002"] == 1
    assert fake_server.detail_request_counts["yield:000003"] == 1

    clock.advance(30)
    finished = crawler.run(run_id, CrawlControl(), lambda event: None)
    output = ExcelExporter(database).export(
        database_scope_id(database),
        tmp_path / "exports" / "anomalies.xlsx",
        None,
    )
    issue_sheet = load_workbook(output, data_only=True)["缺失与异常"]
    issue_values = [
        row[3]
        for row in issue_sheet.iter_rows(min_row=3, values_only=True)
        if row[3] is not None
    ]

    assert finished.success_tasks == 4
    assert finished.failed_tasks == 2
    assert fake_server.detail_request_counts["yield:000001"] == 2
    assert fake_server.detail_request_counts["yield:000002"] == 2
    assert any(row["index_code"] == "000003" and "data date mismatch" in row["last_error"] for row in task_rows(database, run_id))
    assert "数据日期不一致" in issue_values
    with database._connection() as connection:
        malformed = connection.execute(
            """SELECT payload, is_success FROM raw_responses
               WHERE index_code = '000002' AND endpoint = 'yield' AND is_success = 0"""
        ).fetchone()
    assert malformed is not None
    assert malformed["payload"] == "{not-json"
