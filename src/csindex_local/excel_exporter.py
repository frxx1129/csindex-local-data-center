"""Offline Excel export for locally persisted index metrics."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import Database
from .models import MetricRow
from .query_service import METRIC_FIELDS, QueryService


SOURCE_URL = "https://www.csindex.com.cn"
_SHEET_NAMES = ("完整九项指标", "近一年收益率排名", "缺失与异常", "说明")
_METRIC_LABELS = {
    "one_month": "近一月收益率",
    "three_month": "近三月收益率",
    "year_to_date": "年初至今收益率",
    "one_year": "近一年收益率",
    "three_year": "近三年收益率",
    "five_year": "近五年收益率",
    "one_year_volatility": "近一年年化波动率",
    "three_year_volatility": "近三年年化波动率",
    "five_year_volatility": "近五年年化波动率",
}
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=14)
_POSITIVE_FONT = Font(color="C00000")
_NEGATIVE_FONT = Font(color="008000")


class ExcelExporter:
    """Create an Excel workbook entirely from local SQLite data."""

    def __init__(self, database: Database) -> None:
        self._database = database
        self._query = QueryService(database)

    def export(
        self, scope_id: str, output_path: Path, ranking_limit: int | None
    ) -> Path:
        rows = self._query.latest_rows(scope_id)
        ranked = self._query.rank(rows, "one_year", ranking_limit)
        issues = self._query.exception_rows(scope_id, rows)
        generated_at = datetime.now().astimezone().replace(microsecond=0).isoformat()

        workbook = Workbook()
        workbook.remove(workbook.active)
        self._write_complete_sheet(workbook.create_sheet(_SHEET_NAMES[0]), rows)
        self._write_ranking_sheet(workbook.create_sheet(_SHEET_NAMES[1]), ranked)
        self._write_issue_sheet(
            workbook.create_sheet(_SHEET_NAMES[2]), issues, scope_id, rows, generated_at
        )
        self._write_notes_sheet(
            workbook.create_sheet(_SHEET_NAMES[3]), scope_id, rows, generated_at
        )

        return self._save_without_damaging_locked_file(workbook, Path(output_path))

    def _write_complete_sheet(self, sheet, rows: list[MetricRow]) -> None:
        headers = [
            "序号",
            "指数代码",
            "指数名称",
            *(_METRIC_LABELS[field] for field in METRIC_FIELDS),
            "数据日期",
            "缺失项数",
            "完整性",
        ]
        self._start_table(sheet, "完整九项指标", headers)
        for rank, row in enumerate(rows, start=1):
            values = [
                rank,
                row.index_code,
                row.index_name,
                *(
                    value if (value := getattr(row, field)) is not None else "--"
                    for field in METRIC_FIELDS
                ),
                row.data_date,
                row.missing_count,
                "是" if row.is_complete else "否",
            ]
            sheet.append(values)
            self._format_metric_row(sheet, sheet.max_row, range(4, 13))
        self._finish_table(sheet, len(headers))

    def _write_ranking_sheet(self, sheet, rows: list[MetricRow]) -> None:
        headers = ["排名", "指数代码", "指数名称", "近一年收益率", "数据日期"]
        self._start_table(sheet, "近一年收益率排名", headers)
        for rank, row in enumerate(rows, start=1):
            sheet.append([rank, row.index_code, row.index_name, row.one_year, row.data_date])
            self._format_metric_row(sheet, sheet.max_row, (4,))
        self._finish_table(sheet, len(headers))

    def _write_issue_sheet(
        self,
        sheet,
        issues: list[tuple],
        scope_id: str,
        rows: list[MetricRow],
        generated_at: str,
    ) -> None:
        headers = ["指数代码", "指数名称", "数据日期", "类型", "缺失指标", "缺失项数", "说明"]
        self._start_table(sheet, "缺失与异常", headers)
        for code, name, data_date, kind, fields, count, note in issues:
            sheet.append(
                [
                    code,
                    name,
                    data_date or "--",
                    kind,
                    "、".join(_METRIC_LABELS[field] for field in fields) or "--",
                    count if count is not None else "--",
                    note,
                ]
            )
            sheet.cell(sheet.max_row, 1).number_format = "@"
        self._finish_table(sheet, len(headers))
        self._write_metadata(sheet, 9, scope_id, rows, generated_at)

    def _write_notes_sheet(
        self, sheet, scope_id: str, rows: list[MetricRow], generated_at: str
    ) -> None:
        sheet["A1"] = "导出说明"
        sheet["A1"].font = _TITLE_FONT
        self._write_metadata(sheet, 1, scope_id, rows, generated_at, start_row=3)
        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 70
        sheet.freeze_panes = "A3"

    def _write_metadata(
        self,
        sheet,
        start_column: int,
        scope_id: str,
        rows: list[MetricRow],
        generated_at: str,
        *,
        start_row: int = 2,
    ) -> None:
        dates = sorted({row.data_date for row in rows})
        date_range = "、".join(dates) if dates else "--"
        values = (
            ("数据来源", SOURCE_URL),
            ("生成时间", generated_at),
            ("导出范围", f"{scope_id}（{len(rows)} 条最新快照）"),
            ("数据日期", date_range),
            ("缺失口径", "官网空值保留为 NULL，Excel 显示为 --，不估算。"),
        )
        for offset, (label, value) in enumerate(values):
            row = start_row + offset
            label_cell = sheet.cell(row, start_column, label)
            value_cell = sheet.cell(row, start_column + 1, value)
            label_cell.font = Font(bold=True)
            value_cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions[get_column_letter(start_column)].width = 15
        sheet.column_dimensions[get_column_letter(start_column + 1)].width = 30

    @staticmethod
    def _start_table(sheet, title: str, headers: list[str]) -> None:
        sheet["A1"] = title
        sheet["A1"].font = _TITLE_FONT
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(2, column, header)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.freeze_panes = "A3"

    @staticmethod
    def _format_metric_row(sheet, row_number: int, metric_columns) -> None:
        code_cell = sheet.cell(row_number, 2)
        code_cell.number_format = "@"
        for column in metric_columns:
            cell = sheet.cell(row_number, column)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = "0.00"
                if cell.value > 0:
                    cell.font = _POSITIVE_FONT
                elif cell.value < 0:
                    cell.font = _NEGATIVE_FONT

    @staticmethod
    def _finish_table(sheet, column_count: int) -> None:
        sheet.auto_filter.ref = f"A2:{get_column_letter(column_count)}{max(2, sheet.max_row)}"
        for column in range(1, column_count + 1):
            letter = get_column_letter(column)
            longest = max(
                len(str(sheet.cell(row, column).value or ""))
                for row in range(1, sheet.max_row + 1)
            )
            sheet.column_dimensions[letter].width = min(max(longest + 2, 10), 30)
        sheet.row_dimensions[1].height = 24
        sheet.row_dimensions[2].height = 22

    @staticmethod
    def _save_without_damaging_locked_file(workbook: Workbook, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            workbook.save(output_path)
            return output_path
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            candidate = output_path.with_name(
                f"{output_path.stem}_{timestamp}{output_path.suffix}"
            )
            suffix = 1
            while candidate.exists():
                candidate = output_path.with_name(
                    f"{output_path.stem}_{timestamp}_{suffix}{output_path.suffix}"
                )
                suffix += 1
            workbook.save(candidate)
            return candidate
